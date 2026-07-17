"""Last mile: turn the scored feed CSV into the delivered `outcomes_and_satisfaction.xlsx`.

`build_feed` emits the FULL audit surface -- `therapist-scorecard-feed.csv`, ALL evaluated
clinicians (OK + low_volume). The DELIVERABLE is scored-only: this step filters to
`data_quality_flag == 'OK'`, writes a single-sheet `Sheet1` workbook mirroring the live file's
types (integer IDs/volumes as ints, metrics as floats, blank = empty cell / N/A -- NEVER 0), and
(with --deliver) backs up the current live file and atomically replaces it in the ITPowerBiFiles
OneDrive folder.

By default this STAGES + VERIFIES only (writes data/outcomes_and_satisfaction.xlsx, no OneDrive
touch). Pass --deliver to actually overwrite the live file.

Run from repo root:
  python -m evaluation.build_deliverable                # stage + verify (safe)
  python -m evaluation.build_deliverable --deliver      # + overwrite the live OneDrive file
  python -m evaluation.build_deliverable --target PATH  # override the delivery target
"""
from __future__ import annotations
import argparse
import os
import shutil
from datetime import datetime
from pathlib import Path
import pandas as pd

REPO = Path(__file__).resolve().parent.parent
DATA = REPO / "data"
SRC = DATA / "therapist-scorecard-feed.csv"
STAGE = DATA / "outcomes_and_satisfaction.xlsx"
BACKUP_DIR = DATA / "deliverable-backups"   # LOCAL rollback backups (gitignored); NOT the prod folder
KEEP_BACKUPS = 12                           # retain the newest N repo backups (~1 yr of monthly runs)
SHEET = "Sheet1"

DEFAULT_TARGET = Path(
    r"C:\Users\KLM0001\Aegis Therapies, Inc\ITPowerBiFiles - Documents"
    r"\People Dashboard Data Sources\outcomes_and_satisfaction.xlsx"
)

# columns the live file stores as whole integers (blanks -> empty cell). Everything else numeric
# stays float; text stays text. Mirrors the delivered workbook's types exactly.
INT_COLS = [
    "Person_ID", "EmployeeNo", "JobCodeId",
    "Total_Visits", "Total_Minutes", "Total_Tracks",
    "Short_Stay_Tracks", "Short_Stay_Visits", "Short_Stay_Minutes",
    "Long_Stay_Tracks", "Long_Stay_Visits", "Long_Stay_Minutes",
]


def build_frame() -> pd.DataFrame:
    df = pd.read_csv(SRC, encoding="utf-8-sig")   # utf-8-sig: strip the BOM off 'Timeframe'
    n_all = len(df)
    if "data_quality_flag" not in df.columns:
        raise SystemExit("feed CSV has no data_quality_flag column - wrong/old feed?")
    df = df[df["data_quality_flag"] == "OK"].copy()
    print(f"filtered OK: {len(df):,}/{n_all:,} rows (dropped {n_all - len(df):,} low_volume)")

    # integer columns -> nullable Int64 so whole numbers write as ints and blanks stay EMPTY
    # (never coerced to 0). All other numeric columns keep their float NaN = blank cell.
    for c in INT_COLS:
        if c in df.columns:
            df[c] = pd.to_numeric(df[c], errors="coerce").astype("Int64")

    print(f"window: {df['Timeframe'].iloc[0]}  (as_of {df['as_of_date'].iloc[0]})")
    return df


def write_xlsx(df: pd.DataFrame, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(path, engine="openpyxl") as xl:
        df.to_excel(xl, sheet_name=SHEET, index=False)   # NaN/<NA> -> empty cell


def verify(path: Path, expect_rows: int, expect_cols: list[str]) -> None:
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True)
    assert wb.sheetnames == [SHEET], f"sheets {wb.sheetnames} != ['{SHEET}']"
    ws = wb[SHEET]
    hdr = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    assert hdr == expect_cols, "header mismatch vs feed columns"
    assert ws.max_row - 1 == expect_rows, f"rows {ws.max_row - 1} != {expect_rows}"
    wb.close()
    print(f"verified: {SHEET}, {expect_rows:,} data rows, {len(hdr)} cols")


def _prune_backups() -> None:
    """Keep only the newest KEEP_BACKUPS repo backups (they're gitignored and grow unbounded)."""
    baks = sorted(BACKUP_DIR.glob("outcomes_and_satisfaction_*.xlsx"))
    for old in baks[:-KEEP_BACKUPS]:
        old.unlink()
        print(f"pruned old backup -> {old.name}")


def deliver(target: Path, prod_snapshot: bool = False) -> None:
    """Backup the current live file, then atomically replace it with the staged workbook.

    prod_snapshot: also drop a dated `..._backup_YYYY-MM-DD.xlsx` copy of the OUTGOING live file
    INTO the production folder (the March-2026 courtesy) — use ONLY when the schema changed, so the
    IT app-builder can diff old-vs-new. Off by default: routine months leave no prod clutter.
    """
    if not target.parent.exists():
        raise SystemExit(f"target folder missing (OneDrive not synced?): {target.parent}")
    if target.exists():
        BACKUP_DIR.mkdir(parents=True, exist_ok=True)
        stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        bak = BACKUP_DIR / f"{target.stem}_{stamp}{target.suffix}"
        shutil.copy2(target, bak)
        print(f"backed up live file -> {bak}")
        _prune_backups()
        if prod_snapshot:
            snap = target.parent / f"{target.stem}_backup_{datetime.now():%Y-%m-%d}{target.suffix}"
            shutil.copy2(target, snap)
            print(f"prod schema-diff snapshot -> {snap}")
    elif prod_snapshot:
        print("NOTE: --prod-snapshot skipped (no existing live file to snapshot)")
    # write to a temp beside the target (same volume) then atomic replace
    tmp = target.with_suffix(".tmp.xlsx")
    try:
        shutil.copy2(STAGE, tmp)
        os.replace(tmp, target)                 # atomic on same volume
    except PermissionError:
        if tmp.exists():
            tmp.unlink(missing_ok=True)
        raise SystemExit(f"cannot write {target.name} - is it open in Excel or locked by "
                         f"OneDrive sync? Close it and re-run --deliver.")
    print(f"DELIVERED -> {target}")


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--deliver", action="store_true", help="overwrite the live OneDrive file")
    ap.add_argument("--target", type=Path, default=DEFAULT_TARGET, help="delivery target path")
    ap.add_argument("--prod-snapshot", action="store_true",
                    help="on --deliver, also leave a dated backup of the OUTGOING file in the prod "
                         "folder for IT (use ONLY when the schema changed)")
    args = ap.parse_args()

    df = build_frame()
    write_xlsx(df, STAGE)
    verify(STAGE, expect_rows=len(df), expect_cols=list(df.columns))
    print(f"staged -> {STAGE}")

    if args.deliver:
        deliver(args.target, prod_snapshot=args.prod_snapshot)
    else:
        print("stage-only (no OneDrive touch). Pass --deliver to overwrite the live file.")


if __name__ == "__main__":
    main()
